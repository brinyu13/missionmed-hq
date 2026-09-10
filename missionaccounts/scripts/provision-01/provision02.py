#!/usr/bin/env python3
"""Deterministic PROVISION-02 reconciliation using only locally held evidence.

Raw names, emails, WordPress logins, and candidate details are written only to
an owner-readable private evidence directory. Shared outputs are sanitized.
"""
from __future__ import annotations
import argparse, base64, csv, hashlib, json, os, re, stat, unicodedata
from collections import Counter, defaultdict
from pathlib import Path

COURSE_ID=6357
COURSE_NAME='Dr J, Drills On-Call'
SPONSORS={'UCC','MUL'}
EMAIL_RE=re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')

def tokens(value):
    text=unicodedata.normalize('NFKD',str(value or ''))
    text=''.join(ch.lower() if ch.isalnum() and not unicodedata.combining(ch) else ' ' for ch in text)
    return tuple(piece for piece in text.split() if piece)

def normalize_name(value): return ''.join(tokens(value))
def first_last_key(value):
    parts=tokens(value)
    return ''.join((parts[0],parts[-1])) if len(parts)>=2 else ''.join(parts)
def token_set_key(value): return '|'.join(sorted(tokens(value)))
def split_emails(value):
    out=[]
    for piece in re.split(r'[,;\n\r]+',str(value or '')):
        email=piece.strip().lower()
        if EMAIL_RE.fullmatch(email) and email not in out: out.append(email)
    return out
def mask_email(value):
    if not value or '@' not in value:return ''
    local,domain=value.split('@',1);return (local[:1]+'***@'+domain) if local else '***@'+domain

def username_for(first,last,occupied):
    def clean(v):
        s=''.join(ch for ch in unicodedata.normalize('NFKD',str(v or '')) if ch.isascii() and ch.isalnum())
        return s[:1].upper()+s[1:].lower()
    first,last=clean(first),clean(last)
    bases=[first+last[:n] for n in range(1,len(last)+1)] if first and last else [first or last]
    if not bases or not bases[0]:raise ValueError('username_source_missing')
    for candidate in bases:
        if candidate.casefold() not in occupied:return candidate
    i=2
    while f'{bases[-1]}{i}'.casefold() in occupied:i+=1
    return f'{bases[-1]}{i}'

def write_private(path,text):
    path.parent.mkdir(parents=True,exist_ok=True);os.chmod(path.parent,stat.S_IRWXU)
    path.write_text(text,encoding='utf-8');os.chmod(path,stat.S_IRUSR|stat.S_IWUSR)
def json_private(path,value):write_private(path,json.dumps(value,sort_keys=True,indent=2)+'\n')
def json_shared(path,value):
    path.parent.mkdir(parents=True,exist_ok=True);path.write_text(json.dumps(value,sort_keys=True,indent=2)+'\n',encoding='utf-8')

def workbook_rows(path):
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True);ws=wb['Contacts edit'];out=[]
    for number,values in enumerate(ws.iter_rows(min_row=1,max_col=5,values_only=True),1):
        first,middle,last,raw_email,raw_sponsor=values
        name=' '.join(str(x).strip() for x in (first,middle,last) if x and str(x).strip())
        emails=split_emails(raw_email)
        if not name and not emails:continue
        sponsor=str(raw_sponsor or '').strip().upper()
        out.append({'row':number,'first':str(first or '').strip(),'middle':str(middle or '').strip(),'last':str(last or '').strip(),'name':name,'emails':emails,'sponsor':sponsor if sponsor in SPONSORS else 'DIRECT','full_key':normalize_name(name),'first_last_key':normalize_name(first)+normalize_name(last),'token_set_key':token_set_key(name)})
    return out

def build_indexes(rows):
    indexes={k:defaultdict(set) for k in ('full','first_last','token_set','email')}
    for row in rows:
        if row['full_key']:indexes['full'][row['full_key']].add(row['row'])
        if row['first_last_key']:indexes['first_last'][row['first_last_key']].add(row['row'])
        if row['token_set_key']:indexes['token_set'][row['token_set_key']].add(row['row'])
        for email in row['emails']:indexes['email'][email].add(row['row'])
    return indexes

def source_maps(db,supplemental):
    students={str(x['id']):x for x in db['students']};events=defaultdict(list);aliases=defaultdict(list);source_names=defaultdict(set)
    for event in db['events']:events[str(event['student_id'])].append(event)
    for alias in supplemental.get('identity_aliases',[]):aliases[str(alias['student_id'])].append(alias)
    event_owner={str(x['id']):str(x['student_id']) for x in db['events']};sources={str(x['id']):x for x in supplemental.get('source_rows',[])}
    for link in supplemental.get('event_source_links',[]):
        sid=event_owner.get(str(link['attendance_event_id']));source=sources.get(str(link['source_row_id']))
        if sid and source:source_names[sid].add(str(source.get('display_name') or ''))
    return students,events,aliases,source_names

def unique_email_row(user,indexes):
    if not user:return None
    hits=indexes['email'].get(str(user.get('email') or '').strip().lower(),set())
    return next(iter(hits)) if len(hits)==1 else None

def resolve(prior,db,supplemental,wp_candidates,wp_high,rows):
    by_row={x['row']:x for x in rows};idx=build_indexes(rows);students,events,aliases,source_names=source_maps(db,supplemental)
    occupied={str(x).casefold() for x in wp_candidates.get('usernames',[])};output=[]
    for old in prior['rows']:
        if old.get('workbook_row') is not None:continue
        sid=str(old['student_id']);student=students[sid]
        verified_aliases=[a for a in aliases[sid] if a.get('relationship_state')=='verified']
        canonical=str(student.get('display_name') or '')
        event_names={str(x.get('source_display_name') or '') for x in events[sid]}
        known_names={canonical,*event_names,*source_names[sid],*[str(a.get('display_value') or '') for a in verified_aliases]}-{''}
        exact_alias_rows=set()
        for alias in verified_aliases:exact_alias_rows |= idx['full'].get(normalize_name(alias.get('display_value')),set())
        canonical_rows=idx['full'].get(normalize_name(canonical),set())
        attendance_exact=set()
        for name in event_names|source_names[sid]:attendance_exact |= idx['full'].get(normalize_name(name),set())
        first_last_rows=set();token_rows=set();all_exact=set()
        for name in known_names:
            all_exact |= idx['full'].get(normalize_name(name),set())
            first_last_rows |= idx['first_last'].get(first_last_key(name),set())
            token_rows |= idx['token_set'].get(token_set_key(name),set())
        candidate=wp_candidates.get('candidates',{}).get(sid,{})
        exact_wp_rows={r for user in candidate.get('exact_name_matches',[]) for r in [unique_email_row(user,idx)] if r is not None}
        link_wp_rows={r for user in candidate.get('link_matches',[]) for r in [unique_email_row(user,idx)] if r is not None}
        high_sources=[];high_rows=set()
        if len(exact_alias_rows)==1:high_rows |= exact_alias_rows;high_sources.append('exact_verified_identity_alias')
        if len(exact_wp_rows)==1:high_rows |= exact_wp_rows;high_sources.append('existing_wp_email_plus_exact_attendance_identity')
        if len(link_wp_rows)==1:high_rows |= link_wp_rows;high_sources.append('existing_exact_missionaccounts_link_plus_roster_email')
        evidence=[]
        if exact_alias_rows:evidence.append('verified_alias_exact_roster_name')
        if canonical_rows:evidence.append('canonical_exact_roster_name')
        if attendance_exact:evidence.append('effective_attendance_source_exact_roster_name')
        if exact_wp_rows:evidence.append('wp_exact_name_and_roster_email')
        if link_wp_rows:evidence.append('wp_exact_link_and_roster_email')
        if first_last_rows:evidence.append('first_last_candidate')
        if token_rows:evidence.append('token_set_or_reversed_candidate')
        confidence='HOLD';disposition='NO_VALID_ROSTER_MATCH';reason='no_valid_roster_match';workbook_row=None;selected=None
        if len(high_rows)==1:
            workbook_row=by_row[next(iter(high_rows))]
            possible={}
            h=wp_high.get('candidates',{}).get(sid,{})
            for key in ('link_matches','email_matches','name_matches'):
                for user in h.get(key,[]):
                    if str(user.get('missionaccounts_student_id') or '').lower()==sid or str(user.get('email') or '').lower() in workbook_row['emails']:
                        possible[int(user['id'])]=user
            if len(possible)<=1:
                selected=next(iter(possible.values()),None);confidence='HIGH';disposition='RESOLVED_EXISTING_WP' if selected else 'RESOLVED_NEW_WP_NEEDED';reason='+'.join(sorted(set(high_sources)))
            else:
                confidence='REVIEW';disposition='REVIEW_REQUIRED';reason='multiple_wp_candidates_for_high_identity'
        elif len(high_rows)>1:
            confidence='REVIEW';disposition='REVIEW_REQUIRED';reason='conflicting_high_confidence_rows'
        elif all_exact or first_last_rows or token_rows or any(candidate.get(k) for k in ('link_matches','email_matches','exact_name_matches','first_last_matches','token_set_matches')):
            confidence='REVIEW';disposition='REVIEW_REQUIRED';reason='middle_reversed_or_discrepant_identity_requires_review'
        row=workbook_row
        proposed=''
        if row and not selected:
            proposed=username_for(row['first'],row['last'],occupied);occupied.add(proposed.casefold())
        candidate_rows=sorted(all_exact|first_last_rows|token_rows|exact_wp_rows|link_wp_rows)
        candidate_users={int(u['id']):u for key in ('link_matches','email_matches','exact_name_matches','first_last_matches','token_set_matches') for u in candidate.get(key,[])}
        output.append({'student_id':sid,'attendance_name':canonical,'aliases':sorted({str(a.get('display_value') or '') for a in verified_aliases}-{''}),'attendance_months':sorted({str(e['local_day'])[:7] for e in events[sid]}),'attendance_events':len(events[sid]),'candidate_workbook_rows':[{'row':n,'name':by_row[n]['name'],'emails':by_row[n]['emails'],'sponsor':by_row[n]['sponsor']} for n in candidate_rows],'candidate_wp_accounts':[{'id':u['id'],'login':u.get('login'),'email':u.get('email'),'display_name':u.get('display_name'),'missionaccounts_student_id':u.get('missionaccounts_student_id')} for u in candidate_users.values()],'evidence':evidence,'confidence':confidence,'proposed_action':'REUSE_WP_LINK_AND_ENROLL' if selected else ('CREATE_WP_LINK_AND_ENROLL' if row else 'HOLD'),'final_disposition':disposition,'reason':reason,'workbook_row':row['row'] if row else None,'name':row['name'] if row else canonical,'first':row['first'] if row else '','last':row['last'] if row else '','email':row['emails'][0] if row and row['emails'] else '','alternate_emails':row['emails'][1:] if row else [],'sponsor':row['sponsor'] if row else 'UNKNOWN','wp_user_id':selected.get('id') if selected else None,'username':selected.get('login') if selected else '','proposed_username':proposed,'apply':confidence=='HIGH' and row is not None and row['sponsor']=='DIRECT','billing_treatment':'PRESERVE' if not row or row['sponsor']=='DIRECT' else 'HELD_MISSING_DURABLE_SPONSOR_CONTROL'})
    output.sort(key=lambda x:(str(x['attendance_name']).casefold(),x['student_id']))
    return output

def read_created(result_paths):
    found={}
    for path in result_paths:
        if not path.exists() or path.stat().st_size==0:continue
        data=json.load(open(path))
        for row in data.get('operations',[]):
            if row.get('created'):found[int(row['wp_user_id'])]=row
    return found

def write_outputs(args,unmatched,prior,db,rows):
    private=args.private_dir;shared=args.shared_dir;created=read_created(args.auth_results)
    created_sids={str(item.get('student_id') or '') for item in created.values()}
    for row in unmatched:
        if row.get('confidence')=='HIGH':
            row['final_disposition']='RESOLVED_NEW_WP_NEEDED' if row['student_id'] in created_sids else 'RESOLVED_EXISTING_WP'
            row['provision_state']='READY_INVITE_REQUIRED' if row['student_id'] in created_sids else 'READY_EXISTING_LOGIN'
    json_private(private/'unmatched_resolution.json',{'baseline_unmatched':107,'rows':unmatched})
    with (private/'unmatched_resolution.csv').open('w',encoding='utf-8',newline='') as h:
        w=csv.writer(h);w.writerow(['MissionAccounts UUID','Attendance name','Aliases','Attendance months','Candidate workbook rows','Candidate names/emails','Candidate WP accounts','Evidence','Confidence','Proposed action','Final disposition'])
        for r in unmatched:w.writerow([r['student_id'],r['attendance_name'],'; '.join(r['aliases']),'; '.join(r['attendance_months']),'; '.join(str(x['row']) for x in r['candidate_workbook_rows']),'; '.join(x['name']+' '+','.join(x['emails']) for x in r['candidate_workbook_rows']),'; '.join(str(x['id'])+':'+str(x.get('login') or '') for x in r['candidate_wp_accounts']),'; '.join(r['evidence']),r['confidence'],r['proposed_action'],r['final_disposition']])
    os.chmod(private/'unmatched_resolution.csv',stat.S_IRUSR|stat.S_IWUSR)
    updates={r['student_id']:r for r in unmatched};final=[]
    for old in prior['rows']:
        row=updates.get(str(old['student_id']))
        final.append(row if row and row['confidence']=='HIGH' else old)
    json_private(private/'final_reconciliation.json',{'authority':['DR-220','DR-221'],'source_workbook_sha256':hashlib.sha256(args.workbook.read_bytes()).hexdigest(),'course':{'id':COURSE_ID,'name':COURSE_NAME},'rows':final})
    for old in prior['rows']:
        if old.get('action')=='CREATE_WP_LINK_AND_ENROLL' and old.get('wp_user_id'):
            created[int(old['wp_user_id'])]={'wp_user_id':old['wp_user_id'],'student_id':old['student_id'],'email':old.get('email'),'login':old.get('username'),'created':True}
    for row in unmatched:
        if row['confidence']=='HIGH' and row.get('wp_user_id') and int(row['wp_user_id']) in created:
            created[int(row['wp_user_id'])].update({'student_id':row['student_id'],'email':row['email'],'login':row['username']})
    students={str(x['id']):x for x in db['students']};invites=[]
    for uid,item in sorted(created.items()):
        sid=str(item.get('student_id') or '');student=students.get(sid,{})
        invites.append({'name':student.get('display_name') or '','registered_email':item.get('email') or student.get('email') or '','wp_user_id':uid,'username':item.get('login') or '','account_readiness':'READY - INVITE REQUIRED','password_set_reset_eligibility':'ELIGIBLE - NOT SENT','missionaccounts_readiness':'READY' if sid else 'VERIFY','sponsor':'DIRECT','invitation_required':'YES'})
    with (private/'invite_ready_manifest.csv').open('w',encoding='utf-8',newline='') as h:
        fields=['name','registered_email','wp_user_id','username','account_readiness','password_set_reset_eligibility','missionaccounts_readiness','sponsor','invitation_required'];w=csv.DictWriter(h,fieldnames=fields);w.writeheader();w.writerows(invites)
    os.chmod(private/'invite_ready_manifest.csv',stat.S_IRUSR|stat.S_IWUSR)
    with (private/'MX-MISSIONACCOUNTS-5401R-PROVISION-02_FOUNDER.csv').open('w',encoding='utf-8',newline='') as h:
        w=csv.writer(h);w.writerow(['Name','Email','Sponsor','Attendance months','WP status','Username','LearnDash status','MissionAccounts status','Billing state','Invite status','Review note'])
        for r in final:
            high=r.get('confidence')=='HIGH' or r.get('action')=='REUSE_WP_ALREADY_READY'
            w.writerow([r.get('name'),r.get('email'),r.get('sponsor'),'; '.join(r.get('attendance_months',[])),'Existing' if r.get('wp_user_id') else ('Create' if r.get('apply') else 'Held'),r.get('username') or r.get('proposed_username'),'Confirmed' if r.get('course_access') or (high and r.get('wp_user_id')) else ('Add' if r.get('apply') else 'Held'),'Linked' if r.get('current_missionaccounts_link')==r.get('student_id') or (high and r.get('wp_user_id')) else ('Link' if r.get('apply') else 'Held'),r.get('billing_treatment'),'Required' if any(i['wp_user_id']==r.get('wp_user_id') for i in invites) else 'No',r.get('reason')])
    os.chmod(private/'MX-MISSIONACCOUNTS-5401R-PROVISION-02_FOUNDER.csv',stat.S_IRUSR|stat.S_IWUSR)
    sanitized_rows=[{'student_id':r['student_id'],'attendance_months':r['attendance_months'],'candidate_workbook_rows':[x['row'] for x in r['candidate_workbook_rows']],'email_masked':mask_email(r.get('email','')),'sponsor':r['sponsor'],'wp_user_id':r['wp_user_id'],'confidence':r['confidence'],'final_disposition':r['final_disposition'],'reason':r['reason'],'apply':r['apply']} for r in unmatched]
    counts=Counter(r['final_disposition'] for r in unmatched);counts['resolved_high_confidence']=sum(r['confidence']=='HIGH' for r in unmatched);counts['existing_wp_found']=counts['RESOLVED_EXISTING_WP'];counts['new_wp_needed']=counts['RESOLVED_NEW_WP_NEEDED']
    json_shared(shared/'provision02_unmatched_resolution.json',{'baseline_unmatched':107,'counts':dict(counts),'rows':sanitized_rows})
    json_shared(shared/'provision02_reconciliation.json',{'authority':['DR-220','DR-221'],'baseline_commit':'884b02b0f1906e0b2056cfcbdb81c92390021d7a','source_workbook_sha256':hashlib.sha256(args.workbook.read_bytes()).hexdigest(),'course':{'id':COURSE_ID,'name':COURSE_NAME},'counts':dict(counts),'sponsor_exclusion':'BLOCKED_NO_EXISTING_DURABLE_FIELD_AND_NO_DDL_OR_RUNTIME_AUTHORITY','invite_count':len(invites)})
    return dict(counts)|{'invite_count':len(invites),'total_final':len(final)}

def build_wp_snapshot(prior_path,db_path,supplemental_path,workbook_path,template_path,output):
    prior=json.load(open(prior_path));db=json.load(open(db_path));supplemental=json.load(open(supplemental_path));rows=workbook_rows(workbook_path);idx=build_indexes(rows);by_row={r['row']:r for r in rows};students,events,aliases,source_names=source_maps(db,supplemental);candidates=[]
    for old in prior['rows']:
        if old.get('workbook_row') is not None:continue
        sid=str(old['student_id']);names={str(students[sid].get('display_name') or ''),*[str(e.get('source_display_name') or '') for e in events[sid]],*source_names[sid],*[str(a.get('display_value') or '') for a in aliases[sid] if a.get('relationship_state')=='verified']}-{''}
        fulls=sorted({normalize_name(n) for n in names if normalize_name(n)});fls=sorted({first_last_key(n) for n in names if first_last_key(n)});sets=sorted({token_set_key(n) for n in names if token_set_key(n)});candidate_rows=set()
        for key in fulls:candidate_rows |= idx['full'].get(key,set())
        for key in fls:candidate_rows |= idx['first_last'].get(key,set())
        for key in sets:candidate_rows |= idx['token_set'].get(key,set())
        if not candidate_rows:continue
        found_emails={email for number in candidate_rows for email in by_row[number]['emails']};student_email=str(students[sid].get('email') or '').strip().lower()
        if EMAIL_RE.fullmatch(student_email):found_emails.add(student_email)
        candidates.append({'student_id':sid,'emails':sorted(found_emails),'full_name_keys':fulls,'first_last_keys':fls,'token_set_keys':sets})
    payload={'course_id':COURSE_ID,'candidates':candidates};encoded=base64.b64encode(json.dumps(payload,separators=(',',':')).encode()).decode();rendered=template_path.read_text().replace('/*PROVISION02_PAYLOAD*/',f"$provision02_payload=json_decode(base64_decode('{encoded}'),true);");write_private(output,rendered);return len(candidates)

def build_apply(plan_path,template_path,output,mode):
    plan=json.load(open(plan_path));ops=[]
    for r in plan['rows']:
        if r.get('apply'):
            ops.append({'student_id':r['student_id'],'email':r['email'],'display_name':r['name'],'first':r['first'],'last':r['last'],'action':r['proposed_action'],'expected_wp_user_id':r.get('wp_user_id'),'proposed_username':r.get('proposed_username')})
    payload={'mode':mode,'course_id':COURSE_ID,'operations':ops};encoded=base64.b64encode(json.dumps(payload,separators=(',',':')).encode()).decode();text=template_path.read_text().replace('/*PROVISION01_PAYLOAD*/',f"$provision01_payload = json_decode(base64_decode('{encoded}'), true);");write_private(output,text);return len(ops)

def main():
    p=argparse.ArgumentParser();sub=p.add_subparsers(dest='command',required=True)
    r=sub.add_parser('reconcile')
    for name in ('prior','db','supplemental','wp_candidates','wp_high','workbook','private_dir','shared_dir'):r.add_argument('--'+name.replace('_','-'),dest=name,type=Path,required=True)
    r.add_argument('--auth-result',dest='auth_results',type=Path,action='append',default=[])
    w=sub.add_parser('build-wp-snapshot');w.add_argument('--prior',type=Path,required=True);w.add_argument('--db',type=Path,required=True);w.add_argument('--supplemental',type=Path,required=True);w.add_argument('--workbook',type=Path,required=True);w.add_argument('--template',type=Path,required=True);w.add_argument('--output',type=Path,required=True)
    b=sub.add_parser('build-apply');b.add_argument('--plan',type=Path,required=True);b.add_argument('--template',type=Path,required=True);b.add_argument('--output',type=Path,required=True);b.add_argument('--mode',choices=('apply-auth','apply-entitlements'),required=True)
    a=p.parse_args()
    if a.command=='build-wp-snapshot':print(json.dumps({'candidates':build_wp_snapshot(a.prior,a.db,a.supplemental,a.workbook,a.template,a.output)}));return
    if a.command=='build-apply':print(json.dumps({'operations':build_apply(a.plan,a.template,a.output,a.mode)}));return
    data=lambda x:json.load(open(x));prior,db,sup,wp,high=data(a.prior),data(a.db),data(a.supplemental),data(a.wp_candidates),data(a.wp_high);rows=workbook_rows(a.workbook);unmatched=resolve(prior,db,sup,wp,high,rows);print(json.dumps(write_outputs(a,unmatched,prior,db,rows),sort_keys=True))
if __name__=='__main__':main()
