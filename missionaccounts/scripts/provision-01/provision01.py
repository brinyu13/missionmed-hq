#!/usr/bin/env python3
"""Deterministic PROVISION-01 reconciliation with private PII custody.

The script never prints row data. Files containing names or email addresses
must be written below an access-restricted evidence directory.
"""

from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import json
import os
import re
import stat
import unicodedata
import urllib.parse
import urllib.request
from collections import defaultdict
from datetime import date
from pathlib import Path

COURSE_ID = 6357
COURSE_NAME = "Dr J, Drills On-Call"
FROM_DAY = "2026-06-01"
THROUGH_DAY = "2026-08-31"
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
SPONSORS = {"UCC", "MUL"}


def normalize_name(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return "".join(ch.lower() for ch in text if ch.isalnum())


def split_emails(value: object) -> list[str]:
    found: list[str] = []
    for piece in re.split(r"[,;\n\r]+", str(value or "")):
        email = piece.strip().lower()
        if EMAIL_RE.fullmatch(email) and email not in found:
            found.append(email)
    return found


def mask_email(value: str | None) -> str:
    if not value or "@" not in value:
        return ""
    local, domain = value.split("@", 1)
    return (local[:1] + "***@" + domain) if local else "***@" + domain


def username_for(first: str, last: str, occupied: set[str]) -> str:
    first_clean = "".join(ch for ch in unicodedata.normalize("NFKD", first) if ch.isascii() and ch.isalnum())
    last_clean = "".join(ch for ch in unicodedata.normalize("NFKD", last) if ch.isascii() and ch.isalnum())
    first_clean = first_clean[:1].upper() + first_clean[1:].lower()
    last_clean = last_clean[:1].upper() + last_clean[1:].lower()
    if first_clean and last_clean:
        bases = [first_clean + last_clean[:n] for n in range(1, len(last_clean) + 1)]
    elif first_clean or last_clean:
        bases = [first_clean or last_clean]
    else:
        raise ValueError("username_source_missing")
    for candidate in bases:
        if candidate.casefold() not in occupied:
            return candidate
    base = bases[-1]
    suffix = 2
    while f"{base}{suffix}".casefold() in occupied:
        suffix += 1
    return f"{base}{suffix}"


def _write_private(path: Path, data: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    os.chmod(path.parent, stat.S_IRWXU)
    path.write_text(data, encoding="utf-8")
    os.chmod(path, stat.S_IRUSR | stat.S_IWUSR)


def _json_private(path: Path, value: object) -> None:
    _write_private(path, json.dumps(value, sort_keys=True, indent=2) + "\n")


def workbook_rows(path: Path) -> list[dict[str, object]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Contacts edit"]
    rows: list[dict[str, object]] = []
    for row_number, values in enumerate(ws.iter_rows(min_row=1, max_col=5, values_only=True), 1):
        first, middle, last, raw_email, raw_sponsor = values
        emails = split_emails(raw_email)
        full_name = " ".join(str(v).strip() for v in (first, middle, last) if v and str(v).strip())
        sponsor_raw = str(raw_sponsor or "").strip().upper()
        sponsor = sponsor_raw if sponsor_raw in SPONSORS else "DIRECT"
        if not full_name and not emails:
            continue
        rows.append(
            {
                "row": row_number,
                "first": str(first or "").strip(),
                "middle": str(middle or "").strip(),
                "last": str(last or "").strip(),
                "name": full_name,
                "name_norm": normalize_name(full_name),
                "emails": emails,
                "sponsor": sponsor,
            }
        )
    return rows


def _request_json(table: str, params: dict[str, str]) -> list[dict[str, object]]:
    base = os.environ.get("SUPABASE_URL", "").rstrip("/")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "")
    if not base or not key:
        raise RuntimeError("supabase_runtime_credentials_unavailable")
    result: list[dict[str, object]] = []
    start = 0
    while True:
        query = urllib.parse.urlencode(params, safe="(),.*")
        req = urllib.request.Request(
            f"{base}/rest/v1/{table}?{query}",
            headers={
                "apikey": key,
                "Authorization": f"Bearer {key}",
                "Accept-Profile": "missionaccounts",
                "Range": f"{start}-{start + 999}",
            },
        )
        with urllib.request.urlopen(req, timeout=20) as response:
            page = json.loads(response.read().decode("utf-8"))
        if not isinstance(page, list):
            raise RuntimeError("supabase_response_malformed")
        result.extend(page)
        if len(page) < 1000:
            return result
        start += len(page)


def snapshot_db(output: Path) -> dict[str, int]:
    events = _request_json(
        "attendance_event_projection",
        {
            "select": "id,student_id,local_day,cycle_key,source_display_name,interpretation_state,superseded_by_id",
            "interpretation_state": "eq.effective",
            "superseded_by_id": "is.null",
            "and": f"(local_day.gte.{FROM_DAY},local_day.lte.{THROUGH_DAY})",
            "order": "student_id,local_day,id",
        },
    )
    target_ids = sorted({str(e["student_id"]) for e in events if e.get("student_id")})
    projections = _request_json(
        "student_identity_projection",
        {
            "select": "id,display_name,email,joined_at,identity_state,canonical_student_id,absorbed,device_source,excluded",
            "absorbed": "eq.false",
            "device_source": "eq.false",
            "excluded": "eq.false",
            "identity_state": "eq.verified",
            "order": "id",
        },
    )
    by_id = {str(row["id"]): row for row in projections}
    target_ids = [student_id for student_id in target_ids if student_id in by_id]
    target_set = set(target_ids)
    students = [by_id[student_id] for student_id in target_ids]
    links: list[dict[str, object]] = []
    aliases: list[dict[str, object]] = []
    decisions: list[dict[str, object]] = []
    for offset in range(0, len(target_ids), 75):
        batch = target_ids[offset : offset + 75]
        in_filter = f"in.({','.join(batch)})"
        links.extend(_request_json("student", {"select": "id,matrix_user_ref", "id": in_filter, "order": "id"}))
        aliases.extend(
            _request_json(
                "identity_alias_projection",
                {
                    "select": "id,student_id,display_value,relationship_state,confidence,superseded_by_id",
                    "student_id": in_filter,
                    "superseded_by_id": "is.null",
                    "order": "student_id,id",
                },
            )
        )
        decisions.extend(
            _request_json(
                "billing_decision",
                {
                    "select": "id,student_id,cycle_key,treatment,amount_cents,state,superseded_by_id",
                    "student_id": in_filter,
                    "superseded_by_id": "is.null",
                    "order": "student_id,cycle_key,id",
                },
            )
        )
    payload = {
        "generated_on": date.today().isoformat(),
        "window": {"from": FROM_DAY, "through": THROUGH_DAY},
        "course": {"id": COURSE_ID, "name": COURSE_NAME},
        "students": students,
        "links": links,
        "aliases": aliases,
        "events": [e for e in events if str(e.get("student_id")) in target_set],
        "billing_decisions": decisions,
    }
    _json_private(output, payload)
    return {"targets": len(target_ids), "events": len(payload["events"])}


def _workbook_match(student: dict[str, object], aliases: list[str], events: list[dict[str, object]], rows: list[dict[str, object]]) -> tuple[dict[str, object] | None, str, str]:
    known_emails = {str(student.get("email") or "").strip().lower()} - {""}
    email_hits = [row for row in rows if known_emails.intersection(set(row["emails"]))]
    if len(email_hits) == 1:
        return email_hits[0], "HIGH", "exact_verified_email"
    if len(email_hits) > 1:
        return None, "REVIEW", "multiple_workbook_rows_for_email"
    canonical = normalize_name(student.get("display_name"))
    source_names = {normalize_name(e.get("source_display_name")) for e in events}
    alias_names = {normalize_name(a) for a in aliases}
    corroborated = canonical in source_names or canonical in alias_names
    name_hits = [row for row in rows if row["name_norm"] and row["name_norm"] == canonical]
    if len(name_hits) == 1 and corroborated:
        return name_hits[0], "HIGH", "unique_exact_normalized_name_with_attendance_alias"
    if len(name_hits) > 1:
        return None, "REVIEW", "nonunique_workbook_name"
    return None, "REVIEW", "workbook_identity_not_exact"


def build_wp_script(db_path: Path, workbook_path: Path, template_path: Path, output: Path, mode: str, plan_path: Path | None = None) -> dict[str, int]:
    db = json.loads(db_path.read_text(encoding="utf-8"))
    rows = workbook_rows(workbook_path)
    events_by: dict[str, list[dict[str, object]]] = defaultdict(list)
    aliases_by: dict[str, list[str]] = defaultdict(list)
    links = {str(row["id"]): row.get("matrix_user_ref") for row in db["links"]}
    for event in db["events"]:
        events_by[str(event["student_id"])].append(event)
    for alias in db["aliases"]:
        aliases_by[str(alias["student_id"])].append(str(alias.get("display_value") or ""))
    candidates: list[dict[str, object]] = []
    for student in db["students"]:
        sid = str(student["id"])
        match, confidence, reason = _workbook_match(student, aliases_by[sid], events_by[sid], rows)
        emails = list(match["emails"]) if match else []
        if student.get("email") and str(student["email"]).lower() not in emails:
            emails.append(str(student["email"]).lower())
        candidates.append(
            {
                "student_id": sid,
                "display_name": student.get("display_name"),
                "matrix_user_ref": links.get(sid),
                "emails": emails,
                "normalized_names": sorted({normalize_name(student.get("display_name")), *[normalize_name(v) for v in aliases_by[sid]]} - {""}),
                "workbook_confidence": confidence,
                "workbook_reason": reason,
            }
        )
    payload: dict[str, object] = {"mode": mode, "course_id": COURSE_ID, "candidates": candidates}
    if mode in {"apply-auth", "apply-entitlements"}:
        if plan_path is None:
            raise ValueError("apply_plan_required")
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        payload["operations"] = [
            {
                "student_id": row["student_id"],
                "email": row["email"],
                "display_name": row["name"],
                "first": row["first"],
                "last": row["last"],
                "action": row["action"],
                "expected_wp_user_id": row.get("wp_user_id"),
                "proposed_username": row.get("proposed_username"),
            }
            for row in plan["rows"]
            if row.get("apply") is True
        ]
    template = template_path.read_text(encoding="utf-8")
    encoded = base64.b64encode(json.dumps(payload, separators=(",", ":")).encode()).decode()
    rendered = template.replace("/*PROVISION01_PAYLOAD*/", f"$provision01_payload = json_decode(base64_decode('{encoded}'), true);")
    _write_private(output, rendered)
    return {"candidates": len(candidates), "operations": len(payload.get("operations", []))}


def reconcile(db_path: Path, wp_path: Path, workbook_path: Path, private_json: Path, private_csv: Path, sanitized_json: Path) -> dict[str, int]:
    db = json.loads(db_path.read_text(encoding="utf-8"))
    wp = json.loads(wp_path.read_text(encoding="utf-8"))
    rows = workbook_rows(workbook_path)
    events_by: dict[str, list[dict[str, object]]] = defaultdict(list)
    aliases_by: dict[str, list[str]] = defaultdict(list)
    links = {str(row["id"]): row.get("matrix_user_ref") for row in db["links"]}
    for event in db["events"]:
        events_by[str(event["student_id"])].append(event)
    for alias in db["aliases"]:
        aliases_by[str(alias["student_id"])].append(str(alias.get("display_value") or ""))
    occupied = {str(v).casefold() for v in wp.get("usernames", [])}
    output_rows: list[dict[str, object]] = []
    for student in db["students"]:
        sid = str(student["id"])
        workbook_row, confidence, reason = _workbook_match(student, aliases_by[sid], events_by[sid], rows)
        candidate = wp.get("candidates", {}).get(sid, {})
        selected = None
        link_matches = candidate.get("link_matches", [])
        email_matches = candidate.get("email_matches", [])
        if len(link_matches) == 1:
            selected = link_matches[0]
            wp_email = str(selected.get("email") or "").lower()
            if workbook_row and wp_email not in set(workbook_row["emails"]) and wp_email != str(student.get("email") or "").lower():
                confidence, reason = "REVIEW", "linked_wp_email_not_verified"
        elif len(link_matches) > 1:
            confidence, reason = "REVIEW", "duplicate_wp_missionaccounts_links"
        elif len(email_matches) == 1:
            selected = email_matches[0]
        elif len(email_matches) > 1:
            confidence, reason = "REVIEW", "multiple_wp_accounts_for_verified_email"
        if selected and "administrator" in selected.get("roles", []):
            confidence, reason = "REVIEW", "admin_or_drj_excluded"
        sponsor = workbook_row["sponsor"] if workbook_row else "UNKNOWN"
        if sponsor in SPONSORS:
            confidence, reason = "REVIEW", "durable_sponsor_billing_exclusion_not_available"
        email = workbook_row["emails"][0] if workbook_row and workbook_row["emails"] else str(student.get("email") or "").lower()
        action = "REVIEW_IDENTITY"
        proposed_username = ""
        apply = False
        if workbook_row and confidence == "HIGH" and sponsor == "DIRECT":
            if selected:
                linked = selected.get("missionaccounts_student_id") == sid
                enrolled = bool(selected.get("course_access"))
                action = "REUSE_WP_ALREADY_READY" if linked and enrolled else "REUSE_WP_LINK_AND_ENROLL"
            else:
                proposed_username = username_for(str(workbook_row["first"]), str(workbook_row["last"]), occupied)
                occupied.add(proposed_username.casefold())
                action = "CREATE_WP_LINK_AND_ENROLL"
            apply = action != "REUSE_WP_ALREADY_READY"
        months = sorted({str(e["local_day"])[:7] for e in events_by[sid]})
        output_rows.append(
            {
                "student_id": sid,
                "name": str(student.get("display_name") or ""),
                "aliases": sorted({v for v in aliases_by[sid] if v}),
                "first": str(workbook_row["first"]) if workbook_row else "",
                "last": str(workbook_row["last"]) if workbook_row else "",
                "attendance_months": months,
                "workbook_row": workbook_row["row"] if workbook_row else None,
                "email": email,
                "alternate_emails": list(workbook_row["emails"][1:]) if workbook_row else [],
                "sponsor": sponsor,
                "wp_user_id": selected.get("id") if selected else None,
                "username": selected.get("login") if selected else "",
                "wp_email": selected.get("email") if selected else "",
                "current_matrix_user_ref": links.get(sid),
                "current_missionaccounts_link": selected.get("missionaccounts_student_id") if selected else None,
                "course_access": bool(selected.get("course_access")) if selected else False,
                "action": action,
                "proposed_username": proposed_username,
                "billing_treatment": "PRESERVE" if sponsor == "DIRECT" else "HELD_MISSING_DURABLE_SPONSOR_CONTROL",
                "confidence": confidence,
                "reason": reason,
                "apply": apply,
            }
        )
    output_rows.sort(key=lambda row: (str(row["name"]).casefold(), str(row["student_id"])))
    plan = {
        "authority": ["DR-220", "DR-221"],
        "course": {"id": COURSE_ID, "name": COURSE_NAME},
        "source_workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "rows": output_rows,
    }
    _json_private(private_json, plan)
    private_csv.parent.mkdir(parents=True, exist_ok=True)
    os.chmod(private_csv.parent, stat.S_IRWXU)
    with private_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Name", "Email", "Sponsor", "Attendance months", "WP status", "Username", "LearnDash status", "MissionAccounts status", "Billing state", "Invite status", "Review note"])
        for row in output_rows:
            writer.writerow([
                row["name"], row["email"], row["sponsor"], "; ".join(row["attendance_months"]),
                "Existing" if row["wp_user_id"] else ("Create" if row["action"] == "CREATE_WP_LINK_AND_ENROLL" else "Held"),
                row["username"] or row["proposed_username"],
                "Confirmed" if row["course_access"] else ("Add" if row["apply"] else "Held"),
                "Linked" if row["current_missionaccounts_link"] == row["student_id"] else ("Link" if row["apply"] else "Held"),
                row["billing_treatment"],
                "Required" if row["action"] == "CREATE_WP_LINK_AND_ENROLL" else "No",
                row["reason"],
            ])
    os.chmod(private_csv, stat.S_IRUSR | stat.S_IWUSR)
    sanitized = {
        "authority": plan["authority"],
        "course": plan["course"],
        "source_workbook_sha256": plan["source_workbook_sha256"],
        "rows": [
            {
                "student_id": row["student_id"],
                "attendance_months": row["attendance_months"],
                "workbook_row": row["workbook_row"],
                "email_masked": mask_email(str(row["email"])),
                "sponsor": row["sponsor"],
                "wp_user_id": row["wp_user_id"],
                "action": row["action"],
                "billing_treatment": row["billing_treatment"],
                "confidence": row["confidence"],
                "reason": row["reason"],
            }
            for row in output_rows
        ],
    }
    sanitized_json.parent.mkdir(parents=True, exist_ok=True)
    sanitized_json.write_text(json.dumps(sanitized, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    counts: dict[str, int] = defaultdict(int)
    counts["targets"] = len(output_rows)
    for row in output_rows:
        counts[str(row["action"])] += 1
        counts["matched"] += int(row["workbook_row"] is not None)
        counts["apply"] += int(row["apply"] is True)
        counts["direct"] += int(row["sponsor"] == "DIRECT")
        counts["ucc"] += int(row["sponsor"] == "UCC")
        counts["mul"] += int(row["sponsor"] == "MUL")
    return dict(counts)


def main() -> int:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("snapshot-db")
    p.add_argument("--output", type=Path, required=True)
    p = sub.add_parser("build-wp-script")
    p.add_argument("--db", type=Path, required=True)
    p.add_argument("--workbook", type=Path, required=True)
    p.add_argument("--template", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--mode", choices=("snapshot", "apply-auth", "apply-entitlements"), required=True)
    p.add_argument("--plan", type=Path)
    p = sub.add_parser("reconcile")
    p.add_argument("--db", type=Path, required=True)
    p.add_argument("--wp", type=Path, required=True)
    p.add_argument("--workbook", type=Path, required=True)
    p.add_argument("--private-json", type=Path, required=True)
    p.add_argument("--private-csv", type=Path, required=True)
    p.add_argument("--sanitized-json", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "snapshot-db":
        result = snapshot_db(args.output)
    elif args.command == "build-wp-script":
        result = build_wp_script(args.db, args.workbook, args.template, args.output, args.mode, args.plan)
    else:
        result = reconcile(args.db, args.wp, args.workbook, args.private_json, args.private_csv, args.sanitized_json)
    print(json.dumps(result, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
