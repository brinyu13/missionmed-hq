#!/usr/bin/env python3
"""Roster-person-first MissionAccounts PROVISION-03 reconciliation.

Private plans contain identity and email data. Shared output contains only stable
IDs, workbook row numbers, masked email, reasons, and counts.
"""

from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import importlib.util
import json
import os
import re
import stat
import unicodedata
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

COURSE_ID = 6357
COURSE_NAME = "Dr J, Drills On-Call"
SPONSORS = {"UCC", "MUL"}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
STOP_TOKENS = {
    "dr", "doctor", "md", "do", "phd", "ms", "ms3", "ms4", "oms", "omsiii",
    "omsiv", "ii", "iii", "iv", "jr", "sr", "student", "iphone", "ipad",
    "android", "phone", "galaxy", "pixel", "device", "zoom", "computer", "pc",
    "macbook", "tablet", "user", "participant", "guest", "owner", "host", "s",
}
GENERIC_ONLY = {
    "iphone", "ipad", "android", "phone", "galaxy", "pixel", "device", "zoom",
    "computer", "pc", "macbook", "tablet", "user", "participant", "guest", "owner",
    "host", "student", "unknown", "none", "test", "admin", "missionmed",
}


def read_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def write_private(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    os.chmod(path.parent, stat.S_IRWXU)
    text = value if isinstance(value, str) else json.dumps(value, sort_keys=True, indent=2) + "\n"
    path.write_text(text, encoding="utf-8")
    os.chmod(path, stat.S_IRUSR | stat.S_IWUSR)


def write_shared(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, sort_keys=True, indent=2) + "\n", encoding="utf-8")


def ascii_text(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").casefold())
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def name_tokens(value: Any) -> list[str]:
    return [token for token in re.findall(r"[a-z0-9]+", ascii_text(value)) if token not in STOP_TOKENS]


def name_norm(value: Any) -> str:
    return "".join(name_tokens(value))


def similarity(left: str, right: str) -> float:
    if not left or not right:
        return 0.0
    if left == right:
        return 1.0
    return SequenceMatcher(None, left, right).ratio()


def mask_email(value: str) -> str:
    if "@" not in value:
        return ""
    local, domain = value.split("@", 1)
    return (local[:1] + "***@" + domain) if local else "***@" + domain


def split_emails(value: Any) -> list[str]:
    found: list[str] = []
    for item in re.split(r"[\n;,]+", str(value or "")):
        candidate = item.strip().lower()
        if EMAIL_RE.fullmatch(candidate) and candidate not in found:
            found.append(candidate)
    return found


def workbook_rows(path: Path) -> list[dict[str, Any]]:
    source = Path(__file__).with_name("provision02.py")
    if not source.exists():
        source = Path('/Users/brianb/MissionMed_worktrees/MX-MISSIONACCOUNTS-5401R/missionaccounts/scripts/provision-01/provision02.py')
    if source.exists():
        spec = importlib.util.spec_from_file_location("provision02_workbook", source)
        module = importlib.util.module_from_spec(spec)
        assert spec.loader is not None
        spec.loader.exec_module(module)
        rows = module.workbook_rows(path)
        for row in rows:
            row["tokens"] = name_tokens(row["name"])
            row["norm"] = name_norm(row["name"])
        return rows
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Contacts edit"]
    rows = []
    for number, values in enumerate(ws.iter_rows(min_row=1, max_col=5, values_only=True), 1):
        first, middle, last, raw_email, raw_sponsor = values
        name = " ".join(str(x).strip() for x in (first, middle, last) if x and str(x).strip())
        emails = split_emails(raw_email)
        if not name and not emails:
            continue
        sponsor = str(raw_sponsor or "").strip().upper()
        rows.append({"row": number, "first": str(first or "").strip(), "middle": str(middle or "").strip(), "last": str(last or "").strip(), "name": name, "emails": emails, "sponsor": sponsor if sponsor in SPONSORS else "DIRECT", "tokens": name_tokens(name), "norm": name_norm(name)})
    return rows


def username_for(first: str, last: str, occupied: set[str]) -> str:
    first_clean = "".join(ch for ch in first if ch.isalnum())
    last_clean = "".join(ch for ch in last if ch.isalnum())
    bases = [first_clean + last_clean[:n] for n in range(1, len(last_clean) + 1)] if first_clean and last_clean else [first_clean or last_clean]
    if not bases or not bases[0]:
        raise ValueError("username_source_missing")
    for candidate in bases:
        if candidate.casefold() not in occupied:
            return candidate
    suffix = 2
    while f"{bases[-1]}{suffix}".casefold() in occupied:
        suffix += 1
    return f"{bases[-1]}{suffix}"


def best_name_score(source_tokens: list[str], roster_tokens: list[str], unique_first: bool, unique_last: bool) -> tuple[float, str, str]:
    if not source_tokens or not roster_tokens:
        return 0.0, "none", ""
    source_norm = "".join(source_tokens)
    roster_norm = "".join(roster_tokens)
    if source_norm == roster_norm:
        return 1.0, "exact_core", "normalized_full_name_equal"
    if source_tokens == list(reversed(roster_tokens)):
        return 0.995, "reversed", "token_order_reversed"
    if sorted(source_tokens) == sorted(roster_tokens):
        return 0.99, "reversed", "token_set_equal"
    if len(source_tokens) == 1:
        token = source_tokens[0]
        if unique_first and token == roster_tokens[0]:
            return 0.93, "unique_first", "unique_roster_first_name"
        if unique_last and token == roster_tokens[-1]:
            return 0.93, "unique_last", "unique_roster_last_name"
        return max(0.55 * similarity(token, roster_tokens[0]), 0.55 * similarity(token, roster_tokens[-1])), "single_weak", "single_token_not_unique"
    source_first, roster_first = source_tokens[0], roster_tokens[0]
    source_rest, roster_rest = source_tokens[1:], roster_tokens[1:]
    first_score = similarity(source_first, roster_first)
    surname_score = max((similarity(left, right) for left in source_rest for right in roster_rest), default=0.0)
    shared = set(source_rest).intersection(roster_rest)
    if len(source_tokens) >= 2 and set(source_tokens).issubset(roster_tokens):
        return 0.97, "core_tokens", "two_or_more_exact_name_tokens_with_omitted_roster_tokens"
    if source_first == roster_first and shared:
        if set(source_tokens).issubset(roster_tokens) or set(roster_tokens).issubset(source_tokens):
            return 0.98, "core_tokens", "first_name_and_core_surname_with_extra_or_missing_tokens"
        return 0.97, "core_tokens", "first_name_and_surname_root_equal"
    if source_first == roster_first and surname_score >= 0.80:
        return 0.95 if surname_score >= 0.88 else 0.93, "minor_typo", "first_name_equal_and_surname_typo"
    if shared and (
        first_score >= 0.78
        or (min(len(source_first), len(roster_first)) >= 4 and (source_first.startswith(roster_first) or roster_first.startswith(source_first)))
    ):
        return 0.94 if first_score >= 0.88 else 0.92, "minor_typo", "first_name_typo_and_surname_root_equal"
    if source_tokens[-1] == roster_tokens[0] and roster_tokens[-1] == source_tokens[0]:
        return 0.985, "reversed", "first_last_reversed"
    full_score = similarity(source_norm, roster_norm)
    reverse_score = similarity(source_norm, "".join(reversed(roster_tokens)))
    token_score = max((similarity(left, right) for left in source_tokens for right in roster_tokens), default=0.0)
    score = max(full_score, reverse_score) * 0.76 + token_score * 0.24
    return score, "fuzzy", "deterministic_edit_similarity"


def evidence_maps(db: dict[str, Any], supplemental: dict[str, Any]) -> tuple[dict[str, dict[str, Any]], dict[str, set[str]], dict[str, int]]:
    students = {str(row["id"]): row for row in db["students"]}
    names: dict[str, set[str]] = defaultdict(set)
    events: dict[str, int] = Counter()
    for row in db["events"]:
        sid = str(row["student_id"])
        events[sid] += 1
        if row.get("source_display_name"):
            names[sid].add(str(row["source_display_name"]))
    for row in supplemental.get("identity_aliases", []):
        if row.get("superseded_by_id") is None and row.get("relationship_state") in {"verified", "candidate", "device"} and row.get("student_id") and row.get("display_value"):
            names[str(row["student_id"])].add(str(row["display_value"]))
    event_owner = {str(row["id"]): str(row["student_id"]) for row in db["events"]}
    source_rows = {str(row["id"]): row for row in supplemental.get("source_rows", [])}
    for link in supplemental.get("event_source_links", []):
        sid = event_owner.get(str(link.get("attendance_event_id")))
        source = source_rows.get(str(link.get("source_row_id")))
        if sid and source and source.get("display_name"):
            names[sid].add(str(source["display_name"]))
    for sid, student in students.items():
        if student.get("display_name"):
            names[sid].add(str(student["display_name"]))
    return students, names, events


def wp_indexes(snapshot: dict[str, Any]) -> dict[str, Any]:
    indexes: dict[str, Any] = {"users": {}, "email": defaultdict(list), "link": defaultdict(list), "name": defaultdict(list), "usernames": set()}
    for user in snapshot["users"]:
        uid = int(user["id"])
        indexes["users"][uid] = user
        indexes["usernames"].add(str(user.get("login") or "").casefold())
        email = str(user.get("email") or "").strip().lower()
        if email:
            indexes["email"][email].append(user)
        link = str(user.get("missionaccounts_student_id") or "").strip().lower()
        if link:
            indexes["link"][link].append(user)
        for candidate in {user.get("display_name"), " ".join([str(user.get("first_name") or ""), str(user.get("last_name") or "")])}:
            key = name_norm(candidate)
            if key:
                indexes["name"][key].append(user)
    return indexes


def select_wp(row: dict[str, Any], member_ids: list[str], indexes: dict[str, Any]) -> tuple[dict[str, Any] | None, str | None, list[str]]:
    evidence: list[str] = []
    linked = {int(user["id"]): user for sid in member_ids for user in indexes["link"].get(sid.lower(), [])}
    if len(linked) > 1:
        return None, "multiple_independently_linked_wp_accounts", evidence
    if linked:
        user = next(iter(linked.values()))
        evidence.append("existing_exact_missionaccounts_link")
        if "administrator" in user.get("roles", []):
            return None, "administrator_account_conflict", evidence
        if user.get("email") and row["emails"] and str(user["email"]).lower() not in row["emails"]:
            return None, "linked_wp_email_conflicts_with_roster", evidence
        return user, None, evidence
    email_users = {int(user["id"]): user for email in row["emails"] for user in indexes["email"].get(email, [])}
    if len(email_users) > 1:
        return None, "multiple_wp_accounts_for_roster_email", evidence
    if email_users:
        user = next(iter(email_users.values()))
        evidence.append("existing_wp_roster_email")
        if "administrator" in user.get("roles", []):
            return None, "administrator_account_conflict", evidence
        current = str(user.get("missionaccounts_student_id") or "").lower()
        if current and current not in {sid.lower() for sid in member_ids}:
            return None, "wp_account_linked_to_different_student", evidence
        return user, None, evidence
    exact_name_users = {int(user["id"]): user for user in indexes["name"].get(row["norm"], []) if "administrator" not in user.get("roles", [])}
    if len(exact_name_users) == 1:
        user = next(iter(exact_name_users.values()))
        current = str(user.get("missionaccounts_student_id") or "").lower()
        if current and current not in {sid.lower() for sid in member_ids}:
            return None, "wp_name_match_linked_to_different_student", evidence
        evidence.append("existing_wp_unique_exact_name")
        return user, None, evidence
    return None, None, evidence


def analyze(args: argparse.Namespace) -> dict[str, int]:
    prior = read_json(args.prior)
    db = read_json(args.db)
    supplemental = read_json(args.supplemental)
    wp = wp_indexes(read_json(args.wp))
    roster = workbook_rows(args.workbook)
    by_roster = {int(row["row"]): row for row in roster}
    by_email: dict[str, set[int]] = defaultdict(set)
    by_first: dict[str, set[int]] = defaultdict(set)
    by_last: dict[str, set[int]] = defaultdict(set)
    for row in roster:
        for email in row["emails"]:
            by_email[email].add(int(row["row"]))
        if row["tokens"]:
            by_first[row["tokens"][0]].add(int(row["row"]))
            by_last[row["tokens"][-1]].add(int(row["row"]))
    students, names, event_counts = evidence_maps(db, supplemental)
    prior_rows = prior["rows"]
    prior_by_sid = {str(row["student_id"]): row for row in prior_rows}
    roster_to_prior: dict[int, list[str]] = defaultdict(list)
    for row in prior_rows:
        if row.get("workbook_row") is not None:
            roster_to_prior[int(row["workbook_row"])].append(str(row["student_id"]))
    active_same_components: dict[str, set[str]] = {}
    active_different_components: list[set[str]] = []
    for decision in supplemental.get("identity_decisions", []):
        members = {str(value) for value in decision.get("member_student_ids", [])}
        if decision.get("superseded_by_id") is not None:
            continue
        if decision.get("decision") == "same":
            for member in members:
                active_same_components[member] = members
        elif decision.get("decision") == "different":
            active_different_components.append(members)
    known_alias_index: dict[str, set[int]] = defaultdict(set)
    for roster_row, sids in roster_to_prior.items():
        for sid in sids:
            for raw in names[sid]:
                key = name_norm(raw)
                if key and len(key) >= 3:
                    known_alias_index[key].add(roster_row)
    unresolved = [row for row in prior_rows if row.get("workbook_row") is None]
    proposed: dict[str, dict[str, Any]] = {}
    for old in unresolved:
        sid = str(old["student_id"])
        source_names = sorted(names[sid])
        source_email = str(students.get(sid, {}).get("email") or "").strip().lower()
        linked_wp = wp["link"].get(sid.lower(), [])
        decisive_row = None
        decisive_reason = None
        email_rows = by_email.get(source_email, set()) if EMAIL_RE.fullmatch(source_email) else set()
        if len(email_rows) == 1:
            decisive_row = next(iter(email_rows)); decisive_reason = "AUTO_MATCH_EXISTING_WP"
        if len(linked_wp) == 1:
            linked_email = str(linked_wp[0].get("email") or "").lower()
            linked_email_rows = by_email.get(linked_email, set())
            if len(linked_email_rows) == 1:
                decisive_row = next(iter(linked_email_rows)); decisive_reason = "AUTO_MATCH_EXISTING_WP"
        scores = []
        for row in roster:
            best = (0.0, "none", "")
            for raw in source_names:
                st = name_tokens(raw)
                score, rule, detail = best_name_score(st, row["tokens"], len(by_first.get(st[0], set())) == 1 if st else False, len(by_last.get(st[0], set())) == 1 if st else False)
                if score > best[0]:
                    best = (score, rule, detail)
            alias_bonus = 0.0
            for raw in source_names:
                hits = known_alias_index.get(name_norm(raw), set())
                if hits == {int(row["row"])}:
                    alias_bonus = 0.04
                    if best[0] < 0.96:
                        best = (0.96, "known_alias", "unique_existing_alias_chain")
            scores.append((min(1.05, best[0] + alias_bonus), int(row["row"]), best[1], best[2]))
        scores.sort(key=lambda item: (-item[0], item[1]))
        top = scores[0]
        second = scores[1]
        if decisive_row is not None:
            chosen = next(item for item in scores if item[1] == decisive_row)
            top = (max(chosen[0], 1.02), chosen[1], "existing_wp", "unique_verified_wp_or_email_evidence")
            second_score = max(item[0] for item in scores if item[1] != decisive_row)
        else:
            second_score = second[0]
        human_tokens = [token for raw in source_names for token in name_tokens(raw) if token not in GENERIC_ONLY]
        all_initials = bool(human_tokens) and all(len(token) <= 2 for token in human_tokens)
        accepted = False
        if decisive_row is not None:
            accepted = True
        elif top[2] in {"exact_core", "reversed", "core_tokens", "minor_typo", "known_alias"} and top[0] >= 0.92 and (top[0] - second_score >= 0.035 or top[2] in {"exact_core", "reversed", "known_alias"}):
            accepted = True
        elif top[2] in {"unique_first", "unique_last"} and top[0] >= 0.93:
            accepted = True
        elif top[2] == "fuzzy" and top[0] >= 0.88 and top[0] - second_score >= 0.075:
            accepted = True
        prior_target_members = set(roster_to_prior.get(int(top[1]), []))
        same_component = active_same_components.get(sid)
        authority_conflict = None
        if accepted and prior_target_members and same_component and not prior_target_members.intersection(same_component):
            accepted = False
            authority_conflict = "existing_active_identity_merge_component_conflict"
        if accepted and prior_target_members and any(sid in component and prior_target_members.intersection(component) for component in active_different_components):
            accepted = False
            authority_conflict = "existing_distinct_person_decision_conflict"
        if accepted:
            reason = {
                "reversed": "AUTO_MATCH_REVERSED_NAME",
                "unique_first": "AUTO_MATCH_UNIQUE_FIRST_NAME",
                "unique_last": "AUTO_MATCH_UNIQUE_LAST_NAME",
                "known_alias": "AUTO_MATCH_KNOWN_ALIAS",
                "existing_wp": "AUTO_MATCH_EXISTING_WP",
            }.get(top[2], "AUTO_MATCH_UNIQUE_CLOSE_NAME")
            disposition = reason
            workbook_row = top[1]
        else:
            workbook_row = None
            if not human_tokens or (all_initials and top[0] < 0.88) or top[0] < 0.68:
                disposition = "NONHUMAN_OR_INSUFFICIENT_IDENTITY"
            else:
                disposition = "TRUE_AMBIGUITY"
            reason = authority_conflict or disposition
        proposed[sid] = {
            "student_id": sid,
            "original_attendance_identity": str(students.get(sid, {}).get("display_name") or old.get("name") or old.get("attendance_name") or ""),
            "attendance_aliases": source_names,
            "attendance_months": list(old.get("attendance_months", [])),
            "attendance_events": event_counts.get(sid, 0),
            "workbook_row": workbook_row,
            "match_reason": reason,
            "match_evidence": [top[3], f"score={top[0]:.3f}", f"runner_up={second_score:.3f}", f"gap={top[0]-second_score:.3f}"],
            "match_score": round(top[0], 6),
            "runner_up_score": round(second_score, 6),
            "disposition": disposition,
        }
    accepted_groups: dict[int, list[str]] = defaultdict(list)
    for sid, row in proposed.items():
        if row["workbook_row"] is not None:
            accepted_groups[int(row["workbook_row"])].append(sid)
    for roster_row, member_ids in list(accepted_groups.items()):
        components = {tuple(sorted(active_same_components[sid])) for sid in member_ids if sid in active_same_components}
        if not components:
            continue
        if len(components) > 1:
            kept: list[str] = []
        else:
            existing_component = set(next(iter(components)))
            kept = [sid for sid in member_ids if sid in existing_component]
        for sid in set(member_ids) - set(kept):
            proposed[sid]["workbook_row"] = None
            proposed[sid]["disposition"] = "TRUE_AMBIGUITY"
            proposed[sid]["match_reason"] = "existing_active_identity_merge_component_conflict"
        accepted_groups[roster_row] = kept
        if not kept:
            del accepted_groups[roster_row]
    resolved_by_roster: dict[int, list[str]] = defaultdict(list)
    for roster_row, sids in roster_to_prior.items():
        resolved_by_roster[roster_row].extend(sids)
    for roster_row, sids in accepted_groups.items():
        resolved_by_roster[roster_row].extend(sids)
    canonical_by_roster: dict[int, str] = {}
    group_conflicts: dict[int, str] = {}
    for roster_row, member_ids in resolved_by_roster.items():
        unique_members = sorted(set(member_ids))
        prior_members = [sid for sid in unique_members if prior_by_sid[sid].get("workbook_row") is not None]
        linked_members = [sid for sid in unique_members if wp["link"].get(sid.lower())]
        linked_user_ids = {int(user["id"]) for sid in linked_members for user in wp["link"][sid.lower()]}
        if len(linked_user_ids) > 1:
            group_conflicts[roster_row] = "multiple_independently_provisioned_students"
            continue
        if prior_members:
            ready_prior = [sid for sid in prior_members if prior_by_sid[sid].get("wp_user_id")]
            canonical = (ready_prior or prior_members)[0]
        elif linked_members:
            canonical = linked_members[0]
        else:
            canonical = sorted(unique_members, key=lambda sid: (-proposed[sid]["match_score"], -event_counts.get(sid, 0), sid))[0]
        canonical_by_roster[roster_row] = canonical
    final_rows: list[dict[str, Any]] = []
    for old in prior_rows:
        sid = str(old["student_id"])
        p = proposed.get(sid)
        roster_row_number = int(old["workbook_row"]) if old.get("workbook_row") is not None else (int(p["workbook_row"]) if p and p.get("workbook_row") is not None else None)
        if roster_row_number is None or roster_row_number in group_conflicts:
            final_rows.append({
                "student_id": sid,
                "original_attendance_identity": str(students.get(sid, {}).get("display_name") or old.get("name") or ""),
                "resolved_roster_human": "",
                "workbook_row": None,
                "match_reason": group_conflicts.get(roster_row_number, p["match_reason"] if p else "TRUE_AMBIGUITY"),
                "match_evidence": p["match_evidence"] if p else [str(old.get("reason") or "")],
                "disposition": "TRUE_AMBIGUITY" if roster_row_number in group_conflicts else (p["disposition"] if p else "TRUE_AMBIGUITY"),
                "canonical_student_id": None,
                "collapsed_into_already_known_human": False,
                "attendance_months": list(old.get("attendance_months", [])),
                "sponsor": "UNKNOWN",
                "wp_user_id": None,
                "username": "",
                "course_access": False,
                "provision_state": "HELD",
                "was_unresolved": sid in proposed,
            })
            continue
        roster_row = by_roster[roster_row_number]
        canonical = canonical_by_roster[roster_row_number]
        members = sorted(set(resolved_by_roster[roster_row_number]))
        final_rows.append({
            "student_id": sid,
            "original_attendance_identity": str(students.get(sid, {}).get("display_name") or old.get("name") or ""),
            "resolved_roster_human": roster_row["name"],
            "workbook_row": roster_row_number,
            "match_reason": p["match_reason"] if p else "PRESERVED_PROVISION02_MATCH",
            "match_evidence": p["match_evidence"] if p else [str(old.get("reason") or "preserved")],
            "disposition": "ALIAS_OF_ALREADY_MATCHED_STUDENT" if sid != canonical else (p["disposition"] if p else "PRESERVED_MATCH"),
            "canonical_student_id": canonical,
            "collapsed_into_already_known_human": sid != canonical,
            "attendance_months": list(old.get("attendance_months", [])),
            "sponsor": roster_row["sponsor"],
            "wp_user_id": None,
            "username": "",
            "course_access": False,
            "provision_state": "PENDING_PROVIDER_READBACK",
            "was_unresolved": sid in proposed,
        })
    human_groups: list[dict[str, Any]] = []
    occupied = set(wp["usernames"])
    for roster_row_number, canonical in sorted(canonical_by_roster.items()):
        if roster_row_number in group_conflicts:
            continue
        roster_row = by_roster[roster_row_number]
        members = sorted(set(resolved_by_roster[roster_row_number]))
        identity_component_members = sorted(set(members).union(active_same_components.get(canonical, set())))
        user, conflict, wp_evidence = select_wp(roster_row, identity_component_members, wp)
        if conflict:
            provision_action = "HOLD_WP_CONFLICT"
            proposed_username = ""
        elif user:
            provision_action = "REUSE_WP"
            proposed_username = ""
        elif not roster_row["emails"]:
            provision_action = "HOLD_MISSING_EMAIL"
            proposed_username = ""
        else:
            provision_action = "CREATE_WP"
            proposed_username = username_for(roster_row["first"], roster_row["last"], occupied)
            occupied.add(proposed_username.casefold())
        human_groups.append({
            "workbook_row": roster_row_number,
            "name": roster_row["name"],
            "emails": roster_row["emails"],
            "sponsor": roster_row["sponsor"],
            "canonical_student_id": canonical,
            "member_student_ids": members,
            "identity_component_student_ids": identity_component_members,
            "wp_user_id": int(user["id"]) if user else None,
            "username": str(user.get("login") or "") if user else "",
            "wp_email": str(user.get("email") or "") if user else "",
            "current_missionaccounts_link": str(user.get("missionaccounts_student_id") or "") if user else "",
            "course_access": bool(user.get("course_access")) if user else False,
            "provision_action": provision_action,
            "wp_conflict": conflict,
            "proposed_username": proposed_username,
            "wp_evidence": wp_evidence,
            "billing_treatment": "PRESERVE" if roster_row["sponsor"] == "DIRECT" else "SPONSOR_RUNTIME_CONTROL_PENDING",
            "apply_auth": provision_action in {"REUSE_WP", "CREATE_WP"} and (not user or str(user.get("missionaccounts_student_id") or "").lower() != canonical.lower()),
            "apply_entitlement": provision_action in {"REUSE_WP", "CREATE_WP"} and (not user or not bool(user.get("course_access"))),
            "invite_required": provision_action == "CREATE_WP",
        })
    group_by_row = {row["workbook_row"]: row for row in human_groups}
    for row in final_rows:
        group = group_by_row.get(row["workbook_row"])
        if group:
            row["wp_user_id"] = group["wp_user_id"]
            row["username"] = group["username"] or group["proposed_username"]
            row["course_access"] = group["course_access"]
            row["provision_state"] = group["provision_action"]
    plan = {
        "authority": ["DR-220", "DR-221", "FOUNDER_PROVISION03_PRACTICAL_MATCHING_STEER"],
        "source_workbook_sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
        "course": {"id": COURSE_ID, "name": COURSE_NAME},
        "rows": final_rows,
        "humans": human_groups,
    }
    private_dir = args.private_dir
    write_private(private_dir / "provision03_plan.json", plan)
    crosswalk = {
        "authority": plan["authority"],
        "source_workbook_sha256": plan["source_workbook_sha256"],
        "rows": [
            {
                "student_id": row["student_id"], "workbook_row": row["workbook_row"],
                "canonical_student_id": row["canonical_student_id"], "match_reason": row["match_reason"],
                "match_evidence": row["match_evidence"], "sponsor": row["sponsor"],
                "wp_user_id": row["wp_user_id"], "course_access": row["course_access"],
                "disposition": row["disposition"], "collapsed_into_already_known_human": row["collapsed_into_already_known_human"],
            }
            for row in final_rows if str(row["student_id"]) in proposed
        ],
    }
    write_private(private_dir / "provision03_match_crosswalk.json", {**crosswalk, "rows": [dict(item, original_attendance_identity=next(r["original_attendance_identity"] for r in final_rows if r["student_id"] == item["student_id"]), resolved_roster_human=next(r["resolved_roster_human"] for r in final_rows if r["student_id"] == item["student_id"])) for item in crosswalk["rows"]]})
    counts = Counter(row["disposition"] for row in final_rows if row["student_id"] in proposed)
    summary = {
        "prior_unresolved": len(unresolved),
        "counts": dict(counts),
        "resolved_attendance_identities": sum(row["workbook_row"] is not None for row in final_rows),
        "distinct_roster_humans": len(human_groups),
        "aliases_collapsed": sum(bool(row["collapsed_into_already_known_human"]) for row in final_rows),
        "new_roster_humans": len(set(accepted_groups) - set(roster_to_prior)),
        "provider_actions": dict(Counter(row["provision_action"] for row in human_groups)),
        "apply_auth": sum(bool(row["apply_auth"]) for row in human_groups),
        "apply_entitlement": sum(bool(row["apply_entitlement"]) for row in human_groups),
        "sponsored_humans": sum(row["sponsor"] in SPONSORS for row in human_groups),
    }
    write_shared(args.shared_dir / "provision03_reconciliation.json", {"authority": plan["authority"], "course": plan["course"], "source_workbook_sha256": plan["source_workbook_sha256"], **summary})
    print(json.dumps(summary, sort_keys=True))
    return summary


def build_apply(args: argparse.Namespace) -> dict[str, int]:
    plan = read_json(args.plan)
    operations: list[dict[str, Any]] = []
    for human in plan["humans"]:
        if human.get("wp_conflict"):
            continue
        should_apply = bool(human.get("apply_auth")) if args.mode == "apply-auth" else bool(human.get("apply_entitlement"))
        if not should_apply:
            continue
        if args.mode == "apply-entitlement" and not human.get("wp_user_id"):
            raise ValueError("entitlement_target_missing_wp_user")
        operations.append({
            "student_id": human["canonical_student_id"],
            "email": human["emails"][0] if human.get("emails") else "",
            "display_name": human["name"],
            "first": human["name"].split()[0] if human.get("name") else "",
            "last": human["name"].split()[-1] if human.get("name") else "",
            "action": "CREATE_WP_LINK_AND_ENROLL" if human["provision_action"] == "CREATE_WP" else "REUSE_WP_LINK_AND_ENROLL",
            "expected_wp_user_id": human.get("wp_user_id"),
            "proposed_username": human.get("proposed_username") or "",
            "allowed_current_links": human.get("identity_component_student_ids") or human.get("member_student_ids") or [human["canonical_student_id"]],
        })
    operations.sort(key=lambda row: row["student_id"])
    start = args.batch_index * args.batch_size
    selected = operations[start:start + args.batch_size]
    payload = {"mode": args.mode, "course_id": COURSE_ID, "operations": selected}
    encoded = base64.b64encode(json.dumps(payload, separators=(",", ":")).encode()).decode()
    rendered = args.template.read_text(encoding="utf-8").replace("/*PROVISION03_PAYLOAD*/", f"$provision03_payload=json_decode(base64_decode('{encoded}'),true);")
    write_private(args.output, rendered)
    result = {"total_operations": len(operations), "batch_operations": len(selected), "batch_index": args.batch_index}
    print(json.dumps(result, sort_keys=True))
    return result


def update_provider(args: argparse.Namespace) -> dict[str, int]:
    plan = read_json(args.plan)
    humans = {row["canonical_student_id"]: row for row in plan["humans"]}
    auth_operations = []
    entitlement_operations = []
    for path in args.auth_result:
        auth_operations.extend(read_json(path).get("operations", []))
    for path in args.entitlement_result:
        entitlement_operations.extend(read_json(path).get("operations", []))
    for operation in auth_operations:
        human = humans[str(operation["student_id"])]
        human["wp_user_id"] = int(operation["wp_user_id"])
        human["username"] = str(operation.get("login") or "")
        human["wp_email"] = str(operation.get("email") or "")
        human["current_missionaccounts_link"] = str(operation.get("missionaccounts_link") or "")
        human["course_access"] = bool(operation.get("course_access"))
        human["created_in_p03"] = bool(operation.get("created")) or bool(human.get("created_in_p03"))
        human["provision_action"] = "REUSE_WP"
        human["apply_auth"] = False
        human["apply_entitlement"] = not human["course_access"]
    for operation in entitlement_operations:
        human = humans[str(operation["student_id"])]
        if int(operation["wp_user_id"]) != int(human["wp_user_id"]):
            raise ValueError("entitlement_wp_identity_mismatch")
        human["course_access"] = bool(operation.get("course_access"))
        human["apply_entitlement"] = not human["course_access"]
    for row in plan["rows"]:
        human = humans.get(row.get("canonical_student_id"))
        if human:
            row["wp_user_id"] = human.get("wp_user_id")
            row["username"] = human.get("username") or human.get("proposed_username") or ""
            row["course_access"] = bool(human.get("course_access"))
            row["provision_state"] = "READY_INVITE_REQUIRED" if human.get("created_in_p03") else ("READY_EXISTING_LOGIN" if human.get("wp_user_id") and human.get("course_access") and human.get("current_missionaccounts_link", "").lower() == human["canonical_student_id"].lower() else human.get("provision_action"))
    write_private(args.output, plan)
    result = {
        "auth_results": len(auth_operations), "entitlement_results": len(entitlement_operations),
        "ready_humans": sum(bool(row.get("wp_user_id")) and bool(row.get("course_access")) and str(row.get("current_missionaccounts_link") or "").lower() == str(row["canonical_student_id"]).lower() for row in plan["humans"]),
        "created_in_p03": sum(bool(row.get("created_in_p03")) for row in plan["humans"]),
        "pending_auth": sum(bool(row.get("apply_auth")) for row in plan["humans"]),
        "pending_entitlement": sum(bool(row.get("apply_entitlement")) for row in plan["humans"]),
    }
    print(json.dumps(result, sort_keys=True))
    return result


def write_csvs(args: argparse.Namespace) -> dict[str, int]:
    plan = read_json(args.plan)
    output = args.output_dir
    output.mkdir(parents=True, exist_ok=True)
    os.chmod(output, stat.S_IRWXU)
    humans = {row["workbook_row"]: row for row in plan["humans"]}
    crosswalk_rows = [row for row in plan["rows"] if row.get("was_unresolved")]
    crosswalk_json = {
        "authority": plan["authority"], "source_workbook_sha256": plan["source_workbook_sha256"],
        "rows": [
            {
                "original_attendance_identity": row["original_attendance_identity"],
                "resolved_roster_human": row["resolved_roster_human"], "evidence": row["match_evidence"],
                "match_reason": row["match_reason"], "wp_user_id": row.get("wp_user_id"),
                "missionaccounts_uuid": row["student_id"], "canonical_missionaccounts_uuid": row.get("canonical_student_id"),
                "course_access": row.get("course_access"), "disposition": row["disposition"],
                "collapsed_into_already_known_human": row["collapsed_into_already_known_human"],
            } for row in crosswalk_rows
        ],
    }
    write_private(output / "provision03_match_crosswalk.json", crosswalk_json)
    with (output / "provision03_match_crosswalk.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Original attendance identity", "Resolved roster human", "Evidence", "Match reason", "WP user", "MissionAccounts UUID", "Canonical MissionAccounts UUID", "Course access", "Disposition", "Collapsed into already-known human"])
        for row in crosswalk_json["rows"]:
            writer.writerow([row["original_attendance_identity"], row["resolved_roster_human"], "; ".join(row["evidence"]), row["match_reason"], row["wp_user_id"] or "", row["missionaccounts_uuid"], row["canonical_missionaccounts_uuid"] or "", "YES" if row["course_access"] else "NO", row["disposition"], "YES" if row["collapsed_into_already_known_human"] else "NO"])
    invite_rows: list[dict[str, str]] = []
    if args.previous_invites and args.previous_invites.exists():
        with args.previous_invites.open(encoding="utf-8", newline="") as handle:
            invite_rows.extend(dict(row) for row in csv.DictReader(handle))
    existing_wp_ids = {str(row.get("wp_user_id") or "") for row in invite_rows}
    for human in plan["humans"]:
        if not human.get("created_in_p03") or str(human["wp_user_id"]) in existing_wp_ids:
            continue
        invite_rows.append({
            "name": human["name"], "registered_email": human.get("wp_email") or (human["emails"][0] if human.get("emails") else ""),
            "wp_user_id": str(human["wp_user_id"]), "username": human.get("username") or "",
            "account_readiness": "READY - INVITE REQUIRED", "password_set_reset_eligibility": "ELIGIBLE - NOT SENT",
            "missionaccounts_readiness": "READY", "sponsor": human["sponsor"], "invitation_required": "YES",
        })
    invite_fields = ["name", "registered_email", "wp_user_id", "username", "account_readiness", "password_set_reset_eligibility", "missionaccounts_readiness", "sponsor", "invitation_required"]
    with (output / "invite_ready_manifest.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=invite_fields); writer.writeheader(); writer.writerows(invite_rows)
    with (output / "MX-MISSIONACCOUNTS-5401R-PROVISION-03_FOUNDER.csv").open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Attendance identity", "Roster human", "Email", "Sponsor", "Attendance months", "WP status", "Username", "LearnDash status", "MissionAccounts status", "Billing state", "Invite status", "Disposition", "Review note"])
        for row in plan["rows"]:
            human = humans.get(row.get("workbook_row"))
            email = (human.get("wp_email") or (human.get("emails") or [""])[0]) if human else ""
            linked = bool(human and str(human.get("current_missionaccounts_link") or "").lower() == str(human["canonical_student_id"]).lower())
            writer.writerow([row["original_attendance_identity"], row["resolved_roster_human"], email, row["sponsor"], "; ".join(row["attendance_months"]), "Existing" if human and human.get("wp_user_id") else "Held", row.get("username") or "", "Confirmed" if row.get("course_access") else "Held", "Linked" if linked else "Held", human.get("billing_treatment") if human else "UNKNOWN", "Required" if human and human.get("created_in_p03") else "No", row["disposition"], row["match_reason"]])
    for path in (output / "provision03_match_crosswalk.csv", output / "invite_ready_manifest.csv", output / "MX-MISSIONACCOUNTS-5401R-PROVISION-03_FOUNDER.csv"):
        os.chmod(path, stat.S_IRUSR | stat.S_IWUSR)
    result = {"crosswalk_rows": len(crosswalk_rows), "founder_rows": len(plan["rows"]), "invite_rows": len(invite_rows)}
    print(json.dumps(result, sort_keys=True))
    return result


def main() -> int:
    parser = argparse.ArgumentParser()
    sub = parser.add_subparsers(dest="command", required=True)
    run = sub.add_parser("analyze")
    for field in ("prior", "db", "supplemental", "workbook", "wp", "private_dir", "shared_dir"):
        run.add_argument("--" + field.replace("_", "-"), dest=field, type=Path, required=True)
    build = sub.add_parser("build-apply")
    build.add_argument("--plan", type=Path, required=True); build.add_argument("--template", type=Path, required=True)
    build.add_argument("--mode", choices=("apply-auth", "apply-entitlement"), required=True)
    build.add_argument("--batch-index", type=int, required=True); build.add_argument("--batch-size", type=int, default=20)
    build.add_argument("--output", type=Path, required=True)
    update = sub.add_parser("update-provider")
    update.add_argument("--plan", type=Path, required=True); update.add_argument("--auth-result", type=Path, action="append", default=[])
    update.add_argument("--entitlement-result", type=Path, action="append", default=[]); update.add_argument("--output", type=Path, required=True)
    csvs = sub.add_parser("write-csvs")
    csvs.add_argument("--plan", type=Path, required=True); csvs.add_argument("--previous-invites", type=Path)
    csvs.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    if args.command == "analyze":
        analyze(args)
    elif args.command == "build-apply":
        build_apply(args)
    elif args.command == "update-provider":
        update_provider(args)
    elif args.command == "write-csvs":
        write_csvs(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
